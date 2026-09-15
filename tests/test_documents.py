import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import docx
import pymupdf

from errors import InvalidDocumentError, InputLimitError
from parser import DynamicResumeParser


class DocumentTests(unittest.TestCase):
    def setUp(self):
        root = Path(__file__).resolve().parents[1]
        self.tmp = tempfile.TemporaryDirectory(dir=root)
        self.directory = Path(self.tmp.name).resolve()
        assert self.directory.is_relative_to(root)
        self.addCleanup(self.tmp.cleanup)

    def pdf(self, kinds):
        path = self.directory / "resume.pdf"
        with pymupdf.open() as document:
            for kind in kinds:
                page = document.new_page()
                if kind in {"text", "scan_with_text"}:
                    page.insert_text((72, 72), "Candidate Resume: Python engineer with project experience.")
                if kind in {"scan", "scan_with_text"}:
                    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 8, 8), False)
                    pix.clear_with(255)
                    page.insert_image(page.rect, pixmap=pix)
            document.save(path)
        return DynamicResumeParser(str(path))

    def test_text_pdf(self):
        self.assertIn("Python engineer", self.pdf(["text"])._parse_pdf())

    def test_reject_scanned_mixed_blank_and_ocr_layer_before_model(self):
        for kinds in (["scan"], ["text", "scan"], ["scan_with_text"], ["blank"]):
            with self.subTest(kinds=kinds), patch.object(DynamicResumeParser, "_extract_with_llm") as llm:
                with self.assertRaisesRegex(InvalidDocumentError, "authentic"):
                    self.pdf(kinds).process()
                llm.assert_not_called()

    def test_page_limit(self):
        with self.assertRaises(InputLimitError):
            self.pdf(["text"] * 21)._parse_pdf()

    def test_docx_order_headers_footers_and_nested_tables(self):
        path = self.directory / "resume.docx"
        document = docx.Document()
        document.sections[0].header.paragraphs[0].text = "CONTACT"
        document.add_paragraph("FIRST")
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "MIDDLE"
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "NESTED"
        document.add_paragraph("LAST")
        document.sections[0].footer.paragraphs[0].text = "FOOTER"
        document.save(path)
        self.assertEqual(DynamicResumeParser(str(path))._parse_docx().splitlines(),
                         ["CONTACT", "FIRST", "MIDDLE", "NESTED", "LAST", "FOOTER"])

    def test_invalid_documents(self):
        for suffix in ("pdf", "docx"):
            path = self.directory / f"invalid.{suffix}"
            path.write_bytes(b"not a document")
            with self.assertRaises(InvalidDocumentError):
                DynamicResumeParser(str(path)).process()

    def test_empty_docx(self):
        path = self.directory / "empty.docx"
        docx.Document().save(path)
        with self.assertRaises(InvalidDocumentError):
            DynamicResumeParser(str(path)).process()
